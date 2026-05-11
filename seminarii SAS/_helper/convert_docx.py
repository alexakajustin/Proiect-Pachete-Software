"""Convert all seminar docx files to plain text for review."""
import os
import sys

try:
    from docx import Document
except ImportError:
    os.system(f"{sys.executable} -m pip install python-docx")
    from docx import Document

docx_files = [
    (r"G:\salvat-pc-vechi\facultate\Year3\S2\PS\repo\Proiect-Pachete-Software\seminarii SAS\Seminar 1\Seminar 1\Seminar 1\Seminar programare SAS 1 2018.docx", "seminar1.txt"),
    (r"G:\salvat-pc-vechi\facultate\Year3\S2\PS\repo\Proiect-Pachete-Software\seminarii SAS\Seminar 2\Seminar SAS 2\Seminar programare SAS 2.docx", "seminar2.txt"),
    (r"G:\salvat-pc-vechi\facultate\Year3\S2\PS\repo\Proiect-Pachete-Software\seminarii SAS\Seminar 3\Seminar 3\Seminar 3\Seminar programare SAS 3_cu rezolvari.docx", "seminar3.txt"),
    (r"G:\salvat-pc-vechi\facultate\Year3\S2\PS\repo\Proiect-Pachete-Software\seminarii SAS\Seminar 4\Seminar 4\Seminar 4\Seminar programare SAS 4_cu rezolvari.docx", "seminar4.txt"),
    (r"G:\salvat-pc-vechi\facultate\Year3\S2\PS\repo\Proiect-Pachete-Software\seminarii SAS\SAS ML\SAS ML\Seminar Pachete - ML in SAS v3.docx", "seminar_ml.txt"),
]

output_dir = os.path.dirname(os.path.abspath(__file__))

for docx_path, out_name in docx_files:
    print(f"Converting: {os.path.basename(docx_path)}")
    try:
        doc = Document(docx_path)
        text_lines = []
        for para in doc.paragraphs:
            text_lines.append(para.text)
        # Also extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = "\t".join(cell.text for cell in row.cells)
                text_lines.append(row_text)
        
        out_path = os.path.join(output_dir, out_name)
        with open(out_path, "w", encoding="utf-8") as f:
            f.write("\n".join(text_lines))
        print(f"  -> {out_name} ({len(text_lines)} lines)")
    except Exception as e:
        print(f"  ERROR: {e}")

print("\nDone!")
